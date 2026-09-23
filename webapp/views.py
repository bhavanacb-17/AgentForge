from pathlib import Path
import io
import shutil
import zipfile

from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db.models import Q
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from agentforge_core.models import (
    Project,
    ProjectDocumentation,
    ProjectFile,
    ProjectHistory,
    ProjectMessage,
)

from agents.pipeline import AgentPipeline
from core.gemini_client import GeminiClient
from projects.project_storage import ProjectStorage


# ============================================================
# HOME
# ============================================================

def home(request):
    return render(
        request,
        "webapp/home.html",
    )


# ============================================================
# LOGIN
# ============================================================

def login_view(request):

    error = None

    if request.method == "POST":

        username = request.POST.get(
            "username",
            "",
        ).strip()

        password = request.POST.get(
            "password",
            "",
        )

        user = authenticate(
            request,
            username=username,
            password=password,
        )

        if user is not None:

            login(
                request,
                user,
            )

            return redirect("dashboard")

        error = "Invalid username or password."

    return render(
        request,
        "webapp/login.html",
        {
            "error": error,
        },
    )


# ============================================================
# REGISTER
# ============================================================

def register_view(request):

    error = None

    if request.method == "POST":

        username = request.POST.get(
            "username",
            "",
        ).strip()

        email = request.POST.get(
            "email",
            "",
        ).strip()

        password = request.POST.get(
            "password",
            "",
        )

        confirm_password = request.POST.get(
            "confirm_password",
            "",
        )

        if not username or not email or not password:

            error = "All fields are required."

        elif password != confirm_password:

            error = "Passwords do not match."

        elif len(password) < 6:

            error = "Password must contain at least 6 characters."

        elif User.objects.filter(
            username=username
        ).exists():

            error = "Username already exists."

        elif User.objects.filter(
            email=email
        ).exists():

            error = "Email is already registered."

        else:

            user = User.objects.create_user(
                username=username,
                email=email,
                password=password,
            )

            login(
                request,
                user,
            )

            return redirect("dashboard")

    return render(
        request,
        "webapp/register.html",
        {
            "error": error,
        },
    )


# ============================================================
# LOGOUT
# ============================================================

def logout_view(request):

    logout(request)

    return redirect("home")


# ============================================================
# DASHBOARD
# ============================================================

@login_required
def dashboard(request):

    projects = Project.objects.filter(
        owner=request.user
    ).order_by("-updated_at")

    search = request.GET.get(
        "search",
        "",
    ).strip()

    status = request.GET.get(
        "status",
        "",
    ).strip()

    if search:

        projects = projects.filter(
            Q(name__icontains=search)
            |
            Q(description__icontains=search)
        )

    valid_statuses = {
        value
        for value, label in Project.STATUS_CHOICES
    }

    if status in valid_statuses:

        projects = projects.filter(
            status=status
        )

    return render(
        request,
        "webapp/dashboard.html",
        {
            "projects": projects,
            "search": search,
            "selected_status": status,
        },
    )


# ============================================================
# CREATE PROJECT
# ============================================================

@login_required
def create_project(request):

    error = None

    if request.method == "POST":

        name = request.POST.get(
            "name",
            "",
        ).strip()

        description = request.POST.get(
            "description",
            "",
        ).strip()

        requirements = request.POST.get(
            "requirements",
            "",
        ).strip()

        category = request.POST.get(
            "category",
            "",
        ).strip()

        programming_language = request.POST.get(
            "programming_language",
            "",
        ).strip()

        framework = request.POST.get(
            "framework",
            "",
        ).strip()

        database = request.POST.get(
            "database",
            "",
        ).strip()

        additional_instructions = request.POST.get(
            "additional_instructions",
            "",
        ).strip()

        if not name:

            error = "Project name is required."

        elif Project.objects.filter(
            owner=request.user,
            name=name,
        ).exists():

            error = (
                "You already have a project "
                "with this name."
            )

        else:

            project = Project.objects.create(
                name=name,
                description=description,
                requirements=requirements,
                category=category,
                programming_language=programming_language,
                framework=framework,
                database=database,
                additional_instructions=additional_instructions,
                status="active",
                ai_model="Gemini",
                owner=request.user,
            )

            try:

                ProjectStorage().create_project_folder(
                    project.id
                )

            except Exception:

                pass

            ProjectHistory.objects.create(
                project=project,
                event_type="project_created",
                description=(
                    f"Project created by "
                    f"{request.user.username}."
                ),
            )

            return redirect(
                "project-workspace",
                project_id=project.id,
            )

    return render(
        request,
        "webapp/create_project.html",
        {
            "error": error,
        },
    )


# ============================================================
# EDIT PROJECT
# ============================================================

@login_required
def edit_project(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    error = None

    if request.method == "POST":

        name = request.POST.get(
            "name",
            "",
        ).strip()

        duplicate_name = (
            Project.objects
            .filter(
                owner=request.user,
                name=name,
            )
            .exclude(
                id=project.id
            )
            .exists()
        )

        if not name:

            error = "Project name is required."

        elif duplicate_name:

            error = (
                "You already have another "
                "project with this name."
            )

        else:

            project.name = name

            project.description = request.POST.get(
                "description",
                "",
            ).strip()

            project.requirements = request.POST.get(
                "requirements",
                "",
            ).strip()

            project.category = request.POST.get(
                "category",
                "",
            ).strip()

            project.programming_language = request.POST.get(
                "programming_language",
                "",
            ).strip()

            project.framework = request.POST.get(
                "framework",
                "",
            ).strip()

            project.database = request.POST.get(
                "database",
                "",
            ).strip()

            project.additional_instructions = request.POST.get(
                "additional_instructions",
                "",
            ).strip()

            status = request.POST.get(
                "status",
                project.status,
            ).strip()

            valid_statuses = {
                value
                for value, label in Project.STATUS_CHOICES
            }

            if status in valid_statuses:

                project.status = status

            project.save()

            ProjectHistory.objects.create(
                project=project,
                event_type="project_updated",
                description=(
                    f"Project updated by "
                    f"{request.user.username}."
                ),
            )

            return redirect(
                "project-workspace",
                project_id=project.id,
            )

    return render(
        request,
        "webapp/edit_project.html",
        {
            "project": project,
            "error": error,
        },
    )


# ============================================================
# DELETE PROJECT
# ============================================================

@login_required
def delete_project(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    if request.method == "POST":

        project_name = project.name

        try:

            project_folder = Path(
                ProjectStorage().get_project_folder(
                    project.id
                )
            )

            if project_folder.exists():

                shutil.rmtree(
                    project_folder
                )

        except Exception:

            pass

        project.delete()

        return redirect("dashboard")

    return render(
        request,
        "webapp/delete_project.html",
        {
            "project": project,
        },
    )


# ============================================================
# DUPLICATE PROJECT
# ============================================================

@login_required
def duplicate_project(
    request,
    project_id,
):

    source = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    base_name = f"{source.name} Copy"

    new_name = base_name

    counter = 2

    while Project.objects.filter(
        owner=request.user,
        name=new_name,
    ).exists():

        new_name = (
            f"{base_name} {counter}"
        )

        counter += 1

    duplicated = Project.objects.create(
        name=new_name,
        description=source.description,
        requirements=source.requirements,
        category=source.category,
        programming_language=source.programming_language,
        framework=source.framework,
        database=source.database,
        additional_instructions=source.additional_instructions,
        status="active",
        ai_model=source.ai_model,
        owner=request.user,
    )

    try:

        storage = ProjectStorage()

        source_folder = Path(
            storage.get_project_folder(
                source.id
            )
        )

        destination_folder = Path(
            storage.get_project_folder(
                duplicated.id
            )
        )

        if source_folder.exists():

            shutil.copytree(
                source_folder,
                destination_folder,
                dirs_exist_ok=True,
            )

        else:

            storage.create_project_folder(
                duplicated.id
            )

    except Exception:

        pass

    ProjectHistory.objects.create(
        project=duplicated,
        event_type="project_duplicated",
        description=(
            f"Duplicated from project "
            f"'{source.name}'."
        ),
    )

    return redirect(
        "project-workspace",
        project_id=duplicated.id,
    )


# ============================================================
# PROJECT WORKSPACE
# ============================================================

@login_required
def project_workspace(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    return render(
        request,
        "webapp/project_workspace.html",
        {
            "project": project,
        },
    )


# ============================================================
# AI CHAT
# ============================================================

@login_required
def project_chat(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    if request.method == "POST":

        action = request.POST.get(
            "action",
            "",
        ).strip()

        if action == "clear":

            ProjectMessage.objects.filter(
                project=project
            ).delete()

            ProjectHistory.objects.create(
                project=project,
                event_type="chat_cleared",
                description=(
                    "Project chat history "
                    "was cleared."
                ),
            )

            return redirect(
                "project-chat",
                project_id=project.id,
            )

        message = request.POST.get(
            "message",
            "",
        ).strip()

        if message:

            ProjectMessage.objects.create(
                project=project,
                role="user",
                message=message,
            )

            prompt = f"""
You are AgentForge AI.

Project Name:
{project.name}

Description:
{project.description}

Requirements:
{project.requirements}

Category:
{project.category}

Programming Language:
{project.programming_language}

Framework:
{project.framework}

Database:
{project.database}

Additional Instructions:
{project.additional_instructions}

User Message:
{message}

Provide a useful software engineering response.
"""

            try:

                response = GeminiClient().generate(
                    prompt
                )

                if not response:

                    response = (
                        "AgentForge could not "
                        "generate an AI response."
                    )

            except Exception:

                response = (
                    "AgentForge could not connect "
                    "to the configured AI provider."
                )

            ProjectMessage.objects.create(
                project=project,
                role="assistant",
                message=str(response),
            )

            ProjectHistory.objects.create(
                project=project,
                event_type="chat_message_sent",
                description=(
                    "A message was sent in "
                    "the project AI chat."
                ),
            )

        return redirect(
            "project-chat",
            project_id=project.id,
        )

    messages = ProjectMessage.objects.filter(
        project=project
    ).order_by("created_at")

    return render(
        request,
        "webapp/project_chat.html",
        {
            "project": project,
            "messages": messages,
        },
    )


# ============================================================
# REQUIREMENTS
# ============================================================

@login_required
def project_requirements(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    if request.method == "POST":

        project.requirements = request.POST.get(
            "requirements",
            "",
        ).strip()

        project.additional_instructions = request.POST.get(
            "additional_instructions",
            "",
        ).strip()

        project.save()

        ProjectHistory.objects.create(
            project=project,
            event_type="requirements_updated",
            description=(
                "Project requirements were updated."
            ),
        )

        return redirect(
            "project-requirements",
            project_id=project.id,
        )

    return render(
        request,
        "webapp/project_requirements.html",
        {
            "project": project,
        },
    )


# ============================================================
# AGENT PIPELINE
# ============================================================

@login_required
def project_pipeline(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    result = None
    error = None

    if request.method == "POST":

        project_data = {
            "id": project.id,
            "name": project.name,
            "description": project.description,
            "requirements": project.requirements,
            "category": project.category,
            "programming_language": (
                project.programming_language
            ),
            "framework": project.framework,
            "database": project.database,
            "additional_instructions": (
                project.additional_instructions
            ),
        }

        try:

            ProjectHistory.objects.create(
                project=project,
                event_type="pipeline_started",
                description=(
                    "AI agent pipeline started."
                ),
            )

            pipeline = AgentPipeline(
                project_data
            )

            # IMPORTANT:
            # This starts your existing seven-agent
            # AgentForge workflow.
            result = pipeline.run()

            ProjectHistory.objects.create(
                project=project,
                event_type="pipeline_completed",
                description=(
                    "AI agent pipeline completed."
                ),
            )

            # Only mark completed when the pipeline
            # actually returns successfully.
            if isinstance(result, dict):

                pipeline_status = str(
                    result.get(
                        "pipeline_status",
                        "",
                    )
                ).upper()

                if pipeline_status in {
                    "SUCCESS",
                    "COMPLETED",
                    "SUCCESSFUL",
                }:

                    project.status = "completed"

                    project.save()

        except Exception as exc:

            error = str(exc)

            ProjectHistory.objects.create(
                project=project,
                event_type="pipeline_failed",
                description=(
                    f"Pipeline failed: {error}"
                ),
            )

    return render(
        request,
        "webapp/project_pipeline.html",
        {
            "project": project,
            "result": result,
            "error": error,
        },
    )


# ============================================================
# GENERATED FILES
# ============================================================

@login_required
def project_files(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    storage = ProjectStorage()

    project_folder = Path(
        storage.get_project_folder(
            project.id
        )
    )

    files = []

    if project_folder.exists():

        for file_path in sorted(
            project_folder.rglob("*")
        ):

            if file_path.is_file():

                relative_path = (
                    file_path.relative_to(
                        project_folder
                    )
                )

                files.append(
                    {
                        "name": file_path.name,
                        "path": str(
                            relative_path
                        ),
                        "size": file_path.stat().st_size,
                    }
                )

    return render(
        request,
        "webapp/project_files.html",
        {
            "project": project,
            "files": files,
        },
    )


# ============================================================
# SAFE FILE RESOLVER
# ============================================================

def _resolve_project_file(
    project,
    relative_path,
):

    project_folder = Path(
        ProjectStorage().get_project_folder(
            project.id
        )
    ).resolve()

    requested_file = (
        project_folder / relative_path
    ).resolve()

    try:

        requested_file.relative_to(
            project_folder
        )

    except ValueError:

        raise Http404(
            "Invalid project file path."
        )

    if not requested_file.exists():

        raise Http404(
            "Project file does not exist."
        )

    if not requested_file.is_file():

        raise Http404(
            "Requested path is not a file."
        )

    return requested_file


# ============================================================
# DOWNLOAD SINGLE FILE
# ============================================================

@login_required
def download_project_file(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    relative_path = request.GET.get(
        "path",
        "",
    ).strip()

    if not relative_path:

        raise Http404(
            "File path is required."
        )

    file_path = _resolve_project_file(
        project,
        relative_path,
    )

    return FileResponse(
        open(
            file_path,
            "rb",
        ),
        as_attachment=True,
        filename=file_path.name,
    )


# ============================================================
# DOWNLOAD PROJECT ZIP
# ============================================================

@login_required
def download_project_zip(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    project_folder = Path(
        ProjectStorage().get_project_folder(
            project.id
        )
    )

    if not project_folder.exists():

        raise Http404(
            "Project folder does not exist."
        )

    buffer = io.BytesIO()

    with zipfile.ZipFile(
        buffer,
        "w",
        zipfile.ZIP_DEFLATED,
    ) as archive:

        for file_path in project_folder.rglob("*"):

            if file_path.is_file():

                archive.write(
                    file_path,
                    file_path.relative_to(
                        project_folder
                    ),
                )

    buffer.seek(0)

    ProjectHistory.objects.create(
        project=project,
        event_type="project_downloaded",
        description=(
            "Complete project downloaded "
            "as ZIP."
        ),
    )

    filename = (
        project.name.replace(
            " ",
            "_",
        )
        + ".zip"
    )

    return FileResponse(
        buffer,
        as_attachment=True,
        filename=filename,
    )


# ============================================================
# DEFAULT DOCUMENTATION
# ============================================================

def _default_documentation(project):

    content = [
        f"# {project.name}",
        "",
        "## 1. Project Overview",
        "",
        f"Project: {project.name}",
        "",
        project.description or (
            "This project was created using the "
            "AgentForge autonomous multi-agent "
            "software engineering platform."
        ),
        "",
        "AgentForge transforms software requirements "
        "into a structured software project using "
        "specialized AI engineering agents.",
        "",
        "## 2. Problem Statement",
        "",
        "The project is based on the following requirements:",
        "",
        project.requirements or (
            "No project requirements have been "
            "specified yet."
        ),
        "",
        "The purpose of the system is to transform "
        "these requirements into an organized and "
        "maintainable software solution.",
        "",
        "## 3. Objectives",
        "",
        "- Analyze software requirements.",
        "- Design the application architecture.",
        "- Generate application source code.",
        "- Generate tests.",
        "- Detect and repair implementation problems.",
        "- Perform code review.",
        "- Generate technical documentation.",
        "- Maintain project files.",
        "- Provide an integrated AI software "
        "engineering workspace.",
        "",
        "## 4. Features",
        "",
        "### Project Management",
        "",
        "- Create projects",
        "- Edit projects",
        "- Duplicate projects",
        "- Delete projects",
        "- Archive projects",
        "- Restore projects",
        "",
        "### AI Engineering",
        "",
        "- Requirements Analysis",
        "- Architecture Design",
        "- Code Generation",
        "- Test Generation",
        "- Debugging",
        "- Code Review",
        "- Documentation Generation",
        "",
        "### Workspace",
        "",
        "- AI Chat",
        "- Requirements",
        "- Agent Pipeline",
        "- Generated Files",
        "- Documentation",
        "- History",
        "- Settings",
        "",
        "## 5. Agent Pipeline",
        "",
        "AgentForge contains seven specialized agents:",
        "",
        "1. Requirements Analysis",
        "2. Architect",
        "3. Developer",
        "4. Tester",
        "5. Debugger",
        "6. Code Reviewer",
        "7. Documentation",
        "",
        "The agents operate as a software engineering "
        "workflow rather than as an isolated chatbot.",
        "",
        "## 6. Architecture",
        "",
        "User",
        "|",
        "v",
        "Django Web Interface",
        "|",
        "v",
        "Django Backend",
        "|",
        "v",
        "Agent Pipeline",
        "|",
        "+--> Requirements Analysis",
        "|",
        "+--> Architect",
        "|",
        "+--> Developer",
        "|",
        "+--> Tester",
        "|",
        "+--> Debugger",
        "|",
        "+--> Code Reviewer",
        "|",
        "+--> Documentation",
        "|",
        "v",
        "Generated Project",
        "|",
        "+--> Files",
        "+--> Documentation",
        "+--> History",
        "",
        "## 7. Technology Stack",
        "",
        f"Programming Language: "
        f"{project.programming_language or 'Python'}",
        "",
        f"Framework: "
        f"{project.framework or 'Django'}",
        "",
        f"Database: "
        f"{project.database or 'MySQL'}",
        "",
        "API: Django REST Framework",
        "",
        "AI Model: Gemini",
        "",
        "Frontend: Django Templates, HTML, CSS, JavaScript",
        "",
        "Backend: Django",
        "",
        "ORM: Django ORM",
        "",
        "Version Control: Git / GitHub",
        "",
        "## 8. Database",
        "",
        "AgentForge uses MySQL as the persistent database.",
        "",
        "The main application entities include:",
        "",
        "- Users",
        "- Projects",
        "- Project Messages",
        "- Project Files",
        "- Project History",
        "- Project Documentation",
        "",
        "## 9. REST API",
        "",
        "The AgentForge REST API is available under:",
        "",
        "/api/",
        "",
        "Project API endpoints include:",
        "",
        "GET     /api/projects/",
        "POST    /api/projects/",
        "GET     /api/projects/<id>/",
        "PATCH   /api/projects/<id>/",
        "DELETE  /api/projects/<id>/",
        "",
        "Additional endpoints support:",
        "",
        "- Project messages",
        "- Project files",
        "- Project history",
        "- Project documentation",
        "",
        "## 10. Installation",
        "",
        "Install the required dependencies:",
        "",
        "pip install -r requirements.txt",
        "",
        "Apply Django migrations:",
        "",
        "python manage.py migrate",
        "",
        "Create an administrator account if required:",
        "",
        "python manage.py createsuperuser",
        "",
        "Start the AgentForge application:",
        "",
        "python manage.py runserver",
        "",
        "Open the application at:",
        "",
        "http://127.0.0.1:8000/",
        "",
        "## 11. Usage",
        "",
        "1. Register a new account or log in.",
        "2. Open the project dashboard.",
        "3. Create a project.",
        "4. Enter the project requirements.",
        "5. Open the project workspace.",
        "6. Review the requirements.",
        "7. Start the Agent Pipeline.",
        "8. Review the results of the seven agents.",
        "9. Inspect the generated files.",
        "10. Review the generated documentation.",
        "11. Download the generated project if required.",
        "",
        "## 12. Agent Pipeline Details",
        "",
        "### Requirements Analysis",
        "",
        "Analyzes the user requirements and identifies "
        "the functional and technical needs of the project.",
        "",
        "### Architect",
        "",
        "Designs the application architecture and determines "
        "the major components required by the system.",
        "",
        "### Developer",
        "",
        "Generates the project implementation and source "
        "code files.",
        "",
        "### Tester",
        "",
        "Generates and validates tests for the generated "
        "application.",
        "",
        "### Debugger",
        "",
        "Identifies implementation problems and attempts "
        "to repair generated code.",
        "",
        "### Code Reviewer",
        "",
        "Reviews the generated implementation for quality, "
        "consistency, and potential problems.",
        "",
        "### Documentation",
        "",
        "Generates technical documentation describing "
        "the project and its implementation.",
        "",
        "## 13. Validation",
        "",
        "AgentForge validates generated projects after "
        "the agent workflow.",
        "",
        "Validation can include:",
        "",
        "- Python syntax validation",
        "- Framework validation",
        "- Database technology validation",
        "- Generated-file validation",
        "- Test validation",
        "- Final project validation",
        "",
        "## 14. Project History",
        "",
        "AgentForge records important project events including:",
        "",
        "- Project creation",
        "- Project updates",
        "- Requirements updates",
        "- Chat messages",
        "- Pipeline execution",
        "- Pipeline completion",
        "- Documentation generation",
        "- Documentation updates",
        "- Project downloads",
        "- Generated-file deletion",
        "",
        "## 15. AI Integration",
        "",
        "AgentForge uses Gemini as the configured AI provider.",
        "",
        "AI capabilities include:",
        "",
        "- Requirements analysis",
        "- Architecture generation",
        "- Code generation",
        "- Test generation",
        "- Debugging",
        "- Code review",
        "- Documentation generation",
        "- AI project chat",
        "",
        "The application also supports fallback behavior "
        "when the configured AI provider is unavailable.",
        "",
        "## 16. Security",
        "",
        "The Django application provides:",
        "",
        "- User authentication",
        "- Login protection",
        "- CSRF protection",
        "- Project ownership",
        "- Owner-based project access",
        "- Safe project-file resolution",
        "- Protected project operations",
        "- Database-backed persistence",
        "",
        "## 17. Future Enhancements",
        "",
        "Possible future enhancements include:",
        "",
        "- GitHub integration",
        "- Automatic Git commits",
        "- Pull request generation",
        "- AI code explanation",
        "- AI bug fixing",
        "- Project version management",
        "- Celery and Redis",
        "- CI/CD integration",
        "- Docker support",
        "- Cloud deployment",
        "- Multi-user collaboration",
        "- Advanced analytics",
        "- Multiple AI providers",
        "",
        "## 18. Conclusion",
        "",
        f"{project.name} demonstrates an autonomous "
        "software engineering workflow where specialized "
        "AI agents cooperate to transform software "
        "requirements into a structured software project.",
        "",
        "AgentForge combines project management, AI "
        "assistance, automated development, testing, "
        "debugging, code review, documentation, and "
        "generated-file management into a unified "
        "software engineering environment.",
    ]

    return "\n".join(content)


# ============================================================
# DOCUMENTATION
# ============================================================

@login_required
def project_documentation(request, project_id):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    documentation, created = (
        ProjectDocumentation.objects.get_or_create(
            project=project,
            defaults={
                "content": _default_documentation(project)
            },
        )
    )

    if request.method == "POST":

        action = request.POST.get(
            "action",
            "",
        ).strip()

        if action == "save":

            documentation.content = request.POST.get(
                "content",
                "",
            )

            documentation.save()

            ProjectHistory.objects.create(
                project=project,
                event_type="documentation_updated",
                description=(
                    "Project documentation was updated."
                ),
            )

            return redirect(
                "project-documentation",
                project_id=project.id,
            )

        if action == "generate":

            fallback_content = _default_documentation(
                project
            )

            generated_content = fallback_content

            try:

                prompt = f"""
Create professional technical documentation
for the following AgentForge project.

Project Name:
{project.name}

Description:
{project.description}

Requirements:
{project.requirements}

Category:
{project.category}

Programming Language:
{project.programming_language}

Framework:
{project.framework}

Database:
{project.database}

Additional Instructions:
{project.additional_instructions}

Include:

1. Project Overview
2. Problem Statement
3. Objectives
4. Features
5. Agent Pipeline
6. Architecture
7. Technology Stack
8. Database
9. REST API
10. Installation
11. Usage
12. Testing
13. Project History
14. AI Integration
15. Security
16. Future Enhancements
17. Conclusion

Return clean Markdown documentation.
"""

                ai_response = GeminiClient().generate(
                    prompt
                )

                if ai_response:

                    candidate = str(
                        ai_response
                    ).strip()

                    if (
                        candidate
                        and "AgentForge mock response:"
                        not in candidate
                    ):

                        generated_content = candidate

            except Exception:

                generated_content = fallback_content

            documentation.content = generated_content

            documentation.save()

            ProjectHistory.objects.create(
                project=project,
                event_type="documentation_generated",
                description=(
                    "Project documentation was generated."
                ),
            )

            return redirect(
                "project-documentation",
                project_id=project.id,
            )

    return render(
        request,
        "webapp/project_documentation.html",
        {
            "project": project,
            "documentation": documentation,
        },
    )


# ============================================================
# DOCX DOWNLOAD
# ============================================================

@login_required
def download_documentation_docx(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    documentation, created = (
        ProjectDocumentation.objects.get_or_create(
            project=project,
            defaults={
                "content": _default_documentation(project)
            },
        )
    )

    try:
        from docx import Document
    except ImportError:
        return HttpResponse(
            "python-docx is not installed.",
            status=500,
        )

    document = Document()

    document.add_heading(
        project.name,
        level=0,
    )

    content = (
        documentation.content
        or _default_documentation(project)
    )

    for line in content.splitlines():

        stripped = line.strip()

        if not stripped:
            continue

        if stripped.startswith("### "):

            document.add_heading(
                stripped[4:].strip(),
                level=3,
            )

        elif stripped.startswith("## "):

            document.add_heading(
                stripped[3:].strip(),
                level=2,
            )

        elif stripped.startswith("# "):

            document.add_heading(
                stripped[2:].strip(),
                level=1,
            )

        elif stripped.startswith("- "):

            document.add_paragraph(
                stripped[2:].strip(),
                style="List Bullet",
            )

        elif stripped.startswith("```"):

            continue

        else:

            document.add_paragraph(
                stripped
            )

    buffer = io.BytesIO()

    document.save(buffer)

    buffer.seek(0)

    filename = (
        project.name.replace(" ", "_")
        + "_documentation.docx"
    )

    return FileResponse(
        buffer,
        as_attachment=True,
        filename=filename,
    )


# ============================================================
# PDF DOWNLOAD
# ============================================================

@login_required
def download_documentation_pdf(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    documentation, created = (
        ProjectDocumentation.objects.get_or_create(
            project=project,
            defaults={
                "content": _default_documentation(project)
            },
        )
    )

    try:

        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import (
            getSampleStyleSheet,
        )
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )

    except ImportError:

        return HttpResponse(
            "reportlab is not installed.",
            status=500,
        )

    buffer = io.BytesIO()

    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()

    story = [
        Paragraph(
            project.name,
            styles["Title"],
        ),
        Spacer(1, 20),
    ]

    content = (
        documentation.content
        or _default_documentation(project)
    )

    for line in content.splitlines():

        stripped = line.strip()

        if not stripped:

            story.append(
                Spacer(1, 6)
            )

            continue

        safe_text = (
            stripped
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if stripped.startswith("### "):

            story.append(
                Paragraph(
                    safe_text[4:],
                    styles["Heading2"],
                )
            )

        elif stripped.startswith("## "):

            story.append(
                Paragraph(
                    safe_text[3:],
                    styles["Heading1"],
                )
            )

        elif stripped.startswith("# "):

            story.append(
                Paragraph(
                    safe_text[2:],
                    styles["Heading1"],
                )
            )

        elif stripped.startswith("- "):

            story.append(
                Paragraph(
                    "• " + safe_text[2:],
                    styles["BodyText"],
                )
            )

        elif stripped.startswith("```"):

            continue

        else:

            story.append(
                Paragraph(
                    safe_text,
                    styles["BodyText"],
                )
            )

    pdf.build(story)

    buffer.seek(0)

    filename = (
        project.name.replace(" ", "_")
        + "_documentation.pdf"
    )

    return FileResponse(
        buffer,
        as_attachment=True,
        filename=filename,
    )


# ============================================================
# PROJECT HISTORY
# ============================================================

@login_required
def project_history(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    history = ProjectHistory.objects.filter(
        project=project
    ).order_by("-created_at")

    return render(
        request,
        "webapp/project_history.html",
        {
            "project": project,
            "history": history,
        },
    )


# ============================================================
# PROJECT SETTINGS
# ============================================================

@login_required
def project_settings(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    if request.method == "POST":

        action = request.POST.get(
            "action",
            "",
        ).strip()

        if action == "save":

            project.name = request.POST.get(
                "name",
                project.name,
            ).strip()

            project.description = request.POST.get(
                "description",
                project.description,
            ).strip()

            project.programming_language = request.POST.get(
                "programming_language",
                project.programming_language,
            ).strip()

            project.framework = request.POST.get(
                "framework",
                project.framework,
            ).strip()

            project.database = request.POST.get(
                "database",
                project.database,
            ).strip()

            project.ai_model = request.POST.get(
                "ai_model",
                project.ai_model,
            ).strip()

            status = request.POST.get(
                "status",
                project.status,
            ).strip()

            valid_statuses = {
                value
                for value, label in Project.STATUS_CHOICES
            }

            if status in valid_statuses:
                project.status = status

            project.save()

            ProjectHistory.objects.create(
                project=project,
                event_type="project_settings_updated",
                description=(
                    "Project settings were updated."
                ),
            )

            return redirect(
                "project-settings",
                project_id=project.id,
            )

    return render(
        request,
        "webapp/project_settings.html",
        {
            "project": project,
            "status_choices": Project.STATUS_CHOICES,
        },
    )


# ============================================================
# ARCHIVE PROJECT
# ============================================================

@login_required
def archive_project(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    if request.method == "POST":

        project.status = "archived"
        project.save()

        ProjectHistory.objects.create(
            project=project,
            event_type="project_archived",
            description="Project was archived.",
        )

    return redirect(
        "project-settings",
        project_id=project.id,
    )


# ============================================================
# RESTORE PROJECT
# ============================================================

@login_required
def restore_project(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    if request.method == "POST":

        project.status = "active"
        project.save()

        ProjectHistory.objects.create(
            project=project,
            event_type="project_restored",
            description="Project was restored.",
        )

    return redirect(
        "project-settings",
        project_id=project.id,
    )


# ============================================================
# CLEAR PROJECT CHAT
# ============================================================

@login_required
def clear_project_chat(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    if request.method == "POST":

        ProjectMessage.objects.filter(
            project=project
        ).delete()

        ProjectHistory.objects.create(
            project=project,
            event_type="chat_cleared",
            description=(
                "Project chat history was cleared."
            ),
        )

    return redirect(
        "project-settings",
        project_id=project.id,
    )


# ============================================================
# DELETE PROJECT FILES
# ============================================================

@login_required
def delete_project_files(
    request,
    project_id,
):

    project = get_object_or_404(
        Project,
        id=project_id,
        owner=request.user,
    )

    if request.method == "POST":

        project_folder = Path(
            ProjectStorage().get_project_folder(
                project.id
            )
        )

        if project_folder.exists():

            for item in project_folder.iterdir():

                try:

                    if item.is_dir():
                        shutil.rmtree(item)
                    else:
                        item.unlink()

                except Exception:
                    pass

        ProjectFile.objects.filter(
            project=project
        ).delete()

        ProjectHistory.objects.create(
            project=project,
            event_type="project_files_deleted",
            description=(
                "Generated project files were deleted."
            ),
        )

    return redirect(
        "project-settings",
        project_id=project.id,
    )